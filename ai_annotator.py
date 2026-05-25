"""
AI-as-Judge Annotation Script
==============================
Uses GPT-4o to annotate RAG-generated summaries using the same
codebook as human annotators.

Usage:
    pip install openai pandas openpyxl
    python ai_annotator.py

Note: Add your OpenAI API key before running.
"""

import json
import time
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openai import OpenAI

# ─────────────────────────────────────────
# CONFIGURATION
# ─────────────────────────────────────────
OPENAI_API_KEY = "your-openai-api-key-here"   # <-- Add your key
RESULTS_FILE   = "rag_results.xlsx"
OUTPUT_FILE    = "ai_annotation_results.xlsx"
MODEL          = "gpt-4o"

client = OpenAI(api_key=OPENAI_API_KEY)

# ─────────────────────────────────────────
# SYSTEM PROMPT
# ─────────────────────────────────────────
SYSTEM_PROMPT = """You are an expert qualitative researcher specialising in
postcolonial studies and cultural heritage. You will annotate AI-generated
summaries using a structured codebook.

CODEBOOK:
Category 1 — Dominant Narrative (DN):
DN-1: Colonial achievement
DN-2: Neutral historical account
DN-3: Critical/postcolonial
DN-4: Community/indigenous voice
DN-5: Mixed/contested

Category 2 — Absent Perspectives (AP):
AP-1: No obvious absences
AP-2: Colonised community voice absent
AP-3: Decolonial/critical perspective absent
AP-4: Institutional critique absent
AP-5: Multiple perspectives absent
For each absence, assign pipeline stage: S1, S2, or S3

Category 3 — Epistemic Stance (ES):
ES-1: Institutional authority
ES-2: Academic/scholarly authority
ES-3: Distributed authority
ES-4: Community authority
ES-5: Unclear
Also note: Retrieval-driven or Generation-driven

Category 4 — Conceptual Framework (CF):
CF-1: Colonial/institutional language
CF-2: Critical/postcolonial language
CF-3: Neutral/descriptive language
CF-4: Mixed frameworks

Category 5 — Epistemic Contact (EC):
EC-1: Single perspective only
EC-2: Multiple perspectives listed but not engaged
EC-3: Tension acknowledged
EC-4: Critical engagement

Category 6 — Layer 3 Presence (L3):
L3-1: Present and prominent
L3-2: Present but marginal
L3-3/S1: Absent — not in corpus
L3-3/S2: Absent — in corpus but not retrieved
L3-3/S3: Absent — retrieved but not reflected in summary

You must respond ONLY with a valid JSON object."""

# ─────────────────────────────────────────
# ANNOTATION PROMPT
# ─────────────────────────────────────────
def build_prompt(query, summary, retrieved_titles, stances, layer3_count, layer3_present, k):
    return f"""Annotate the following RAG-generated summary using the codebook.

QUERY: {query}
RETRIEVAL CONTEXT (k={k}):
- Retrieved document titles: {retrieved_titles}
- Epistemic stances: {stances}
- Layer 3 document count: {layer3_count}
- Layer 3 present: {layer3_present}

GENERATED SUMMARY:
{summary}

Respond with ONLY this JSON structure:
{{
  "cat1_dominant_narrative": {{"code": "DN-X", "key_phrase": "..."}},
  "cat2_absent_perspectives": {{"codes": ["AP-X"], "stage": "S1/S2/S3", "notes": "..."}},
  "cat3_epistemic_stance": {{"code": "ES-X", "attribution": "Retrieval-driven or Generation-driven", "evidence": "..."}},
  "cat4_conceptual_framework": {{"code": "CF-X", "key_terms": "..."}},
  "cat5_epistemic_contact": {{"code": "EC-X", "notes": "..."}},
  "cat6_layer3_presence": {{"code": "L3-1/L3-2/L3-3/S1/L3-3/S2/L3-3/S3", "evidence": "..."}},
  "qualitative_notes": "2-3 sentences on pipeline epistemic loss",
  "llm_homogenisation_flag": "Yes or No",
  "llm_homogenisation_explanation": "if yes, explain"
}}"""

# ─────────────────────────────────────────
# ANNOTATE ONE SUMMARY
# ─────────────────────────────────────────
def annotate_summary(query, summary, retrieved_titles, stances, layer3_count, layer3_present, k):
    prompt = build_prompt(query, summary, retrieved_titles, stances, layer3_count, layer3_present, k)
    try:
        response = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user",   "content": prompt}
            ],
            temperature=0.1,
            max_tokens=800,
        )
        raw = response.choices[0].message.content.strip()
        raw = raw.replace("```json", "").replace("```", "").strip()
        return json.loads(raw)
    except Exception as e:
        print(f"  Error: {e}")
        return None

# ─────────────────────────────────────────
# RUN ALL ANNOTATIONS
# ─────────────────────────────────────────
def run_all_annotations(df):
    results = []
    total   = len(df)
    for idx, row in df.iterrows():
        query      = row["Query"]
        k          = row["k Value"]
        summary    = str(row["Generated Summary"])
        titles     = str(row["Retrieved Titles (top 5)"])
        stances    = str(row["Stances"])
        l3_count   = int(row["Layer 3 Count"])
        l3_present = str(row["Layer 3 Present?"])

        print(f"[{idx+1}/{total}] k={k}: {query[:50]}...")
        annotation = annotate_summary(query, summary, titles, stances, l3_count, l3_present, k)

        if annotation:
            result = {
                "Query":            query,
                "k Value":          k,
                "L3 Count":         l3_count,
                "L3 Present?":      l3_present,
                "Embedding Variance": row["Embedding Variance"],
                "Cat1 Code":        annotation.get("cat1_dominant_narrative", {}).get("code", ""),
                "Cat1 Key Phrase":  annotation.get("cat1_dominant_narrative", {}).get("key_phrase", ""),
                "Cat2 Code(s)":     "; ".join(annotation.get("cat2_absent_perspectives", {}).get("codes", [])),
                "Cat2 Stage":       annotation.get("cat2_absent_perspectives", {}).get("stage", ""),
                "Cat2 Notes":       annotation.get("cat2_absent_perspectives", {}).get("notes", ""),
                "Cat3 Code":        annotation.get("cat3_epistemic_stance", {}).get("code", ""),
                "Cat3 Attribution": annotation.get("cat3_epistemic_stance", {}).get("attribution", ""),
                "Cat3 Evidence":    annotation.get("cat3_epistemic_stance", {}).get("evidence", ""),
                "Cat4 Code":        annotation.get("cat4_conceptual_framework", {}).get("code", ""),
                "Cat4 Terms":       annotation.get("cat4_conceptual_framework", {}).get("key_terms", ""),
                "Cat5 Code":        annotation.get("cat5_epistemic_contact", {}).get("code", ""),
                "Cat5 Notes":       annotation.get("cat5_epistemic_contact", {}).get("notes", ""),
                "Cat6 Code":        annotation.get("cat6_layer3_presence", {}).get("code", ""),
                "Cat6 Evidence":    annotation.get("cat6_layer3_presence", {}).get("evidence", ""),
                "Qualitative Notes": annotation.get("qualitative_notes", ""),
                "LLM Flag":         annotation.get("llm_homogenisation_flag", ""),
                "LLM Explanation":  annotation.get("llm_homogenisation_explanation", ""),
            }
            results.append(result)
            print(f"  DN:{result['Cat1 Code']} | ES:{result['Cat3 Code']} | L3:{result['Cat6 Code']}")
        else:
            print(f"  Annotation failed")
        time.sleep(1)
    return results

# ─────────────────────────────────────────
# SAVE TO EXCEL
# ─────────────────────────────────────────
DARK_BLUE = "1F3864"
K_COLORS  = {3: "FCE4D6", 10: "DDEBF7", 20: "E2EFDA"}


def save_results(results, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI Annotations"
    ws.freeze_panes = "A2"

    thin   = Side(style="thin", color="AAAAAA")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = list(results[0].keys()) if results else []
    for ci, header in enumerate(headers, 1):
        cell = ws.cell(1, ci, header)
        cell.font      = Font(bold=True, color="FFFFFF", size=10, name="Arial")
        cell.fill      = PatternFill("solid", fgColor=DARK_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = 20
    ws.row_dimensions[1].height = 22

    for ri, result in enumerate(results, 2):
        k        = result.get("k Value", 3)
        row_fill = PatternFill("solid", fgColor=K_COLORS.get(k, "FFFFFF"))
        for ci, header in enumerate(headers, 1):
            cell = ws.cell(ri, ci, result.get(header, ""))
            cell.font      = Font(size=9, name="Arial")
            cell.fill      = row_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = BORDER
        ws.row_dimensions[ri].height = 45

    wb.save(filename)
    print(f"\nSaved: {filename} ({len(results)} annotations)")


if __name__ == "__main__":
    if OPENAI_API_KEY == "your-openai-api-key-here":
        print("Please add your OpenAI API key first.")
    else:
        print("Starting AI-as-Judge Annotation\n")
        df      = pd.read_excel(RESULTS_FILE, sheet_name=0)
        print(f"Loaded {len(df)} rows ({df['Query'].nunique()} queries x 3 k-values)\n")
        results = run_all_annotations(df)
        save_results(results, OUTPUT_FILE)
        print("Done!")
