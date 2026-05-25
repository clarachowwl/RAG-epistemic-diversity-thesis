"""
Layer 2 Corpus Builder — Europeana
====================================
Fetches cross-institutional records from the Europeana API.

Usage:
    pip install requests openpyxl
    python layer2_europeana.py

API key: eefetropi (public research key)
"""

import requests
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

# ─────────────────────────────────────────
# CONFIGURATION
# ─────────────────────────────────────────
API_KEY     = "eefetropi"
OUTPUT_FILE = "layer2_europeana_corpus.xlsx"
MAX_PER_QUERY = 40
BASE_URL    = "https://api.europeana.eu/record/v2/search.json"

SEARCH_QUERIES = [
    ("Indonesia Dutch colonial",          "Indonesia / Dutch East Indies"),
    ("Suriname slavery Dutch",            "Suriname / Slavery"),
    ("VOC Dutch East India Company",      "VOC / Trade"),
    ("Java Batavia Netherlands",          "Java / Batavia"),
    ("Dutch colonial portrait",           "Colonial portraiture"),
]

# ─────────────────────────────────────────
# FETCH FUNCTIONS
# ─────────────────────────────────────────
def search_europeana(query, rows=40, start=1):
    params = [
        ("wskey",   API_KEY),
        ("query",   query),
        ("rows",    rows),
        ("start",   start),
        ("profile", "rich"),
    ]
    r = requests.get(BASE_URL, params=params, timeout=15)
    if r.status_code != 200:
        print(f"  Error {r.status_code}")
        return []
    data = r.json()
    print(f"  API totalResults: {data.get('totalResults', 0)}")
    return data.get("items", [])


def extract_fields(item, topic_label):
    description = ""
    for field in ["dcDescription", "dcDescriptionLangAware"]:
        val = item.get(field)
        if isinstance(val, list) and val:
            description = val[0] if isinstance(val[0], str) else str(val[0])
        elif isinstance(val, dict):
            en = val.get("en", [])
            description = en[0] if en else ""
        if description and len(description.strip()) > 30:
            break

    if not description or len(description.strip()) < 30:
        title_val   = item.get("title", [])
        description = title_val[0] if title_val else ""
    if not description or len(description.strip()) < 10:
        return None

    title_list = item.get("title", [])
    title      = title_list[0] if title_list else ""

    dates        = item.get("year", item.get("dcDate", []))
    date_created = dates[0] if dates else ""

    data_provider = item.get("dataProvider", ["Unknown"])
    provider      = data_provider[0] if isinstance(data_provider, list) else str(data_provider)

    country     = item.get("country", ["Unknown"])
    country_str = country[0] if isinstance(country, list) else str(country)

    creators = item.get("dcCreator", [])
    creator  = creators[0] if creators else "Unknown"

    subjects     = item.get("dcSubject", [])
    if isinstance(subjects, dict):
        subjects = subjects.get("en", [])
    subject_tags = "; ".join(subjects[:5]) if subjects else ""

    rights_list = item.get("rights", [""])
    rights      = rights_list[0] if rights_list else ""

    shown_at   = item.get("edmIsShownAt", [])
    source_url = shown_at[0] if shown_at else item.get("guid", "")

    europeana_id = item.get("id", "").replace("/", "_").strip("_")

    return {
        "doc_id":              f"euro_{europeana_id}",
        "title":               title,
        "date_created":        str(date_created),
        "providing_institution": provider,
        "country_of_provider": country_str,
        "creator":             creator,
        "subject_tags":        subject_tags,
        "topic":               topic_label,
        "description":         description.strip(),
        "record_type":         item.get("type", ""),
        "rights_statement":    rights,
        "language":            "en",
        "source_institution":  f"Europeana / {provider}",
        "epistemic_layer":     2,
        "epistemic_stance":    "institutional",
        "source_url":          source_url,
        "retrieved_date":      str(date.today()),
    }

# ─────────────────────────────────────────
# MAIN FETCH
# ─────────────────────────────────────────
def fetch_corpus():
    records  = []
    seen_ids = set()
    for query, topic_label in SEARCH_QUERIES:
        print(f"\nSearching: '{query}'")
        items = search_europeana(query, rows=MAX_PER_QUERY)
        print(f"  Returned {len(items)} items")
        for item in items:
            item_id = item.get("id", "")
            if item_id in seen_ids:
                continue
            seen_ids.add(item_id)
            record = extract_fields(item, topic_label)
            if record:
                records.append(record)
                print(f"  + {record['providing_institution'][:25]}: {record['title'][:40]}")
            else:
                print(f"  - No usable text, skipped")
            time.sleep(0.1)
        time.sleep(0.3)
    print(f"\nTotal Layer 2 records: {len(records)}")
    return records

# ─────────────────────────────────────────
# EXCEL OUTPUT
# ─────────────────────────────────────────
COLUMNS = [
    ("doc_id",               "Doc ID",              20),
    ("title",                "Title",               35),
    ("date_created",         "Date Created",        14),
    ("providing_institution","Providing Institution",28),
    ("country_of_provider",  "Country",             14),
    ("creator",              "Creator",             22),
    ("subject_tags",         "Subject Tags",        28),
    ("topic",                "Topic",               24),
    ("description",          "Description (EN)",    60),
    ("record_type",          "Record Type",         12),
    ("rights_statement",     "Rights",              20),
    ("source_institution",   "Source Institution",  28),
    ("epistemic_layer",      "Epistemic Layer",     14),
    ("epistemic_stance",     "Epistemic Stance",    18),
    ("source_url",           "Source URL",          40),
    ("retrieved_date",       "Retrieved Date",      16),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")
thin        = Side(style="thin", color="AAAAAA")
BORDER      = Border(left=thin, right=thin, top=thin, bottom=thin)


def save_to_excel(records, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Layer 2 Europeana Corpus"
    ws.freeze_panes = "A2"

    for ci, (field, label, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, ci, label)
        cell.font      = Font(bold=True, color="FFFFFF", size=10, name="Arial")
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 22

    for ri, record in enumerate(records, 2):
        alt = (ri % 2 == 0)
        for ci, (field, label, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(ri, ci, record.get(field, ""))
            cell.font      = Font(size=10, name="Arial")
            cell.fill      = ALT_FILL if alt else PatternFill()
            cell.alignment = Alignment(vertical="top", wrap_text=(field == "description"))
            cell.border    = BORDER
        ws.row_dimensions[ri].height = 55 if len(str(record.get("description", ""))) > 80 else 28

    wb.save(filename)
    print(f"\nSaved: {filename} ({len(records)} records)")


if __name__ == "__main__":
    records = fetch_corpus()
    if records:
        save_to_excel(records, OUTPUT_FILE)
    else:
        print("No records fetched.")
